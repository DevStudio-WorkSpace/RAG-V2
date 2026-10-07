"""
core.sublitex
-------------
Contrato unico de datos de Sublitex. Es la unica fuente de verdad para:

  - Formato y generacion de codigos: SBX-NN-NNNN  (^SBX-\\d{2}-\\d{4}$)
    Definido por SECCION 00 / Tarea 1 ("El cargador"). Estructura:
        SBX   = prefijo fijo
        NN    = numero de carpeta de origen (01..99)
        NNNN  = correlativo dentro de esa carpeta (0001..9999)
    Ejemplo oficial: SBX-03-0147.
  - Planilla "Control de exportacion" con 9 columnas:
    codigo, carpeta_origen, archivo_original, tipo, deporte, ocasion,
    estado, quien, fecha
  - Hoja "Listas": numero de carpeta -> carpeta real del Drive. El numero
    de carpeta tambien esta dentro del codigo y debe coincidir con la
    columna carpeta_origen de la planilla.
  - Trazabilidad obligatoria:
    imagen publicada -> codigo -> carpeta_origen + archivo_original -> CDR

Cualquier otro modulo del proyecto que maneje codigos SBX, la planilla o el
mapa de carpetas debe importar las funciones de aca (unificar, no duplicar).

Regla de seguridad: este modulo NUNCA renombra, mueve, borra ni sobrescribe
CDR. Solo reconstruye su ubicacion a partir de carpeta_origen y
archivo_original. Para trabajar con una copia segura usar core.cdr_seguridad.
"""

import csv
import os
import re

PATRON_CODIGO = r"^SBX-\d{2}-\d{4}$"
REGEX_CODIGO = re.compile(PATRON_CODIGO)

COLUMNAS_PLANILLA = [
    "codigo",
    "carpeta_origen",
    "archivo_original",
    "tipo",
    "deporte",
    "ocasion",
    "estado",
    "quien",
    "fecha",
]
COLUMNAS_TRAZABILIDAD = ["codigo", "carpeta_origen", "archivo_original"]
COLUMNAS_PLANILLA_DISENO = ["tipo", "deporte", "ocasion", "estado"]

HOJA_PRINCIPAL = "Control de exportación"
HOJA_LISTAS = "Listas"


class ErrorCodigo(ValueError):
    """Codigo SBX que no respeta el formato SBX-NN-NNNN."""


class ErrorTrazabilidad(ValueError):
    """No se puede reconstruir la cadena hasta el CDR."""


def validar_codigo(codigo):
    """True si el codigo cumple exactamente SBX-NN-NNNN (carpeta 2 digitos,
    correlativo 4 digitos)."""
    return bool(REGEX_CODIGO.match(str(codigo or "").strip()))


def descomponer_codigo(codigo):
    """Devuelve (numero_de_carpeta, correlativo) validando el formato.

    Para SBX-03-0147 devuelve (3, 147). Lanza ErrorCodigo si no cumple el
    patron de SECCION 00.
    """
    codigo = str(codigo or "").strip()
    if not REGEX_CODIGO.match(codigo):
        raise ErrorCodigo(
            f"Codigo invalido: '{codigo}' (se espera SBX-NN-NNNN)."
        )
    carpeta, correlativo = codigo.split("-")[1:]
    return int(carpeta), int(correlativo)


def generar_codigo(carpeta, correlativo):
    """Genera un codigo SBX-NN-NNNN (SECCION 00 / Tarea 1).

    - carpet >= 1, <= 99 (NN, 2 digitos)
    - correlativo >= 1, <= 9999 (NNNN, 4 digitos dentro de la carpeta)

    El numero de carpeta del codigo DEBE coincidir con la columna
    carpeta_origen de la planilla (se valida en validar_registro).
    """
    try:
        carpeta_int = int(carpeta)
        correlativo_int = int(correlativo)
    except (TypeError, ValueError):
        raise ErrorCodigo("Carpeta y correlativo deben ser enteros.")
    if not (1 <= carpeta_int <= 99):
        raise ErrorCodigo(
            f"El numero de carpeta debe ser 1..99, se recibio {carpeta_int}."
        )
    if not (1 <= correlativo_int <= 9999):
        raise ErrorCodigo(
            f"El correlativo debe ser 1..9999, se recibio {correlativo_int}."
        )
    return "SBX-{:02d}-{:04d}".format(carpeta_int, correlativo_int)


def _entradas(mapa):
    if not mapa:
        return {}
    if "entradas" in mapa:
        return mapa["entradas"]
    return mapa


def _normalizar_numero(texto):
    numero = str(texto or "").strip()
    if numero.isdigit() and len(numero) == 1:
        return numero.zfill(2)
    return numero


def _filtrar_encabezado(pares):
    if not pares:
        return pares
    primera = [str(c or "").strip().lower() for c in pares[0]]
    claves = ("numero", "número", "nro", "carpeta", "nombre", "id", "ruta")
    if any(any(k in c for k in claves) for c in primera):
        return pares[1:]
    return pares


def _mapa_desde_pares(pares):
    entradas = {}
    repetidos = []
    vacios = []
    for indice, par in enumerate(pares, start=1):
        numero = "" if not par else _normalizar_numero(par[0])
        carpeta = "" if len(par) < 2 else str(par[1] or "").strip()
        if numero == "" and carpeta == "":
            continue
        if carpeta == "":
            vacios.append((numero, indice))
        if numero in entradas:
            repetidos.append(numero)
        entradas[numero] = carpeta
    return {"entradas": entradas, "repetidos": repetidos, "vacios": vacios}


def _leer_csv(ruta):
    ultimo_error = None
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            with open(ruta, "r", encoding=encoding, newline="") as f:
                lector = csv.DictReader(f)
                columnas = list(lector.fieldnames or [])
                filas = [dict(fila) for fila in lector]
            return columnas, filas, "CSV ({})".format(encoding)
        except UnicodeDecodeError as error:
            ultimo_error = error
    raise ultimo_error


def leer_planilla(ruta):
    """Lee la planilla "Control de exportacion" (CSV o XLSX).

    Devuelve:
        {"columnas": [...], "filas": [{columna: valor}], "formato": str,
         "hojas": [...]}

    Para XLSX usa la hoja principal (la llamada "Control de exportación" si
    existe; si no, la primera hoja). No modifica el archivo.
    """
    extension = os.path.splitext(ruta)[1].lower()
    if extension not in (".xlsx", ".xlsm"):
        columnas, filas, formato = _leer_csv(ruta)
        return {
            "columnas": columnas,
            "filas": filas,
            "formato": formato,
            "hojas": [],
        }

    from openpyxl import load_workbook

    libro = load_workbook(ruta, read_only=True, data_only=True)
    if HOJA_PRINCIPAL in libro.sheetnames:
        hoja = libro[HOJA_PRINCIPAL]
    else:
        hoja = libro.active
    filas_crudas = [list(fila) for fila in hoja.iter_rows(values_only=True)]
    hojas = list(libro.sheetnames)
    libro.close()

    if not filas_crudas:
        return {"columnas": [], "filas": [], "formato": "XLSX (hoja vacia)", "hojas": hojas}
    columnas = ["" if c is None else str(c).strip() for c in filas_crudas[0]]
    filas = []
    for fila in filas_crudas[1:]:
        if all(c is None or str(c).strip() == "" for c in fila):
            continue
        filas.append({
            columnas[i]: ("" if i >= len(fila) or fila[i] is None else str(fila[i]).strip())
            for i in range(len(columnas))
        })
    return {
        "columnas": columnas,
        "filas": filas,
        "formato": "XLSX (hoja '{}')".format(hoja.title),
        "hojas": hojas,
    }


def _pares_hoja_listas(libro, hoja):
    if hoja not in libro.sheetnames:
        disponibles = ", ".join(libro.sheetnames) or "ninguna"
        libro.close()
        raise ErrorTrazabilidad(
            "La planilla no tiene la hoja '{}'. Hojas disponibles: {}.".format(
                hoja, disponibles
            )
        )
    hoja_listas = libro[hoja]
    pares = []
    for fila in hoja_listas.iter_rows(values_only=True):
        if all(c is None or str(c).strip() == "" for c in fila):
            continue
        pares.append([("" if c is None else str(c).strip()) for c in fila])
    libro.close()
    return _filtrar_encabezado(pares)


def _pares_mapa_csv(ruta):
    columnas, filas, _ = _leer_csv(ruta)

    def indice_de(claves):
        for i, c in enumerate(columnas):
            if str(c or "").strip().lower() in claves:
                return i
        return None

    i_numero = indice_de(("numero", "número", "nro", "id", "carpeta_nro"))
    i_carpeta = indice_de(("carpeta", "carpeta_real", "ruta", "nombre", "carpeta_drive"))
    if i_numero is None:
        i_numero = 0
    if i_carpeta is None:
        i_carpeta = 1
    pares = []
    for fila in filas:
        numero = ""
        carpeta = ""
        if i_numero < len(columnas) and columnas[i_numero] in fila:
            numero = str(fila.get(columnas[i_numero], "") or "").strip()
        if i_carpeta < len(columnas) and columnas[i_carpeta] in fila:
            carpeta = str(fila.get(columnas[i_carpeta], "") or "").strip()
        if numero == "" and carpeta == "":
            continue
        pares.append([numero, carpeta])
    return pares


def leer_mapa(ruta_planilla, hoja=HOJA_LISTAS, ruta_mapa_csv=None):
    """Mapa numero de carpeta -> carpeta real del Drive.

    Fuentes (en orden):
      1. --ruta_mapa_csv: CSV con columnas tipo {numero, carpeta}.
      2. Hoja "Listas" de una planilla XLSX.

    Devuelve:
        {"entradas": {numero: carpeta}, "repetidos": [...], "vacios": [...],
         "origen": str}

    Lanza ErrorTrazabilidad si no hay ninguna fuente utilizable. No inventa
    valores: las carpetas reales deben estar en la hoja Listas.
    """
    if ruta_mapa_csv:
        if not os.path.isfile(ruta_mapa_csv):
            raise ErrorTrazabilidad("No se encontro el CSV del mapa: {}".format(ruta_mapa_csv))
        mapa = _mapa_desde_pares(_pares_mapa_csv(ruta_mapa_csv))
        mapa["origen"] = "CSV {}".format(os.path.basename(ruta_mapa_csv))
        return mapa

    if ruta_planilla and os.path.splitext(ruta_planilla)[1].lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        libro = load_workbook(ruta_planilla, read_only=True, data_only=True)
        pares = _pares_hoja_listas(libro, hoja)
        mapa = _mapa_desde_pares(pares)
        mapa["origen"] = "hoja '{}'".format(hoja)
        return mapa

    raise ErrorTrazabilidad(
        "El mapa de carpetas requiere una planilla XLSX con la hoja '{}' "
        "o un CSV de mapa (--mapa).".format(hoja)
    )


def carpeta_real(mapa, numero):
    """Carpeta real del Drive para un numero de carpeta, segun la hoja Listas."""
    return _entradas(mapa).get(_normalizar_numero(numero), "")


def _numero_de_carpeta(fila):
    """Numero de carpeta de un registro: viene de la columna carpeta_origen.

    En SBX-NN-NNNN el numero de carpeta del codigo tambien indica la carpeta;
    validar_registro exige que ambos coincidan (codigo y planilla).
    """
    return _normalizar_numero(str(fila.get("carpeta_origen", "") or "").strip())


def validar_registro(fila, mapa):
    """Errores de un registro de la planilla frente al contrato de datos.

    Devuelve una lista (vacia si el registro es completo). Un registro es
    valido solo si tiene codigo SBX-NN-NNNN, carpeta_origen y archivo_original,
    la carpeta del codigo coincide con la columna carpeta_origen, y esa carpeta
    existe en el mapa de Listas.
    """
    errores = []
    codigo = str(fila.get("codigo", "") or "").strip()
    if not codigo:
        errores.append("codigo vacio")
        return errores
    if not validar_codigo(codigo):
        errores.append("codigo invalido (se espera SBX-NN-NNNN): {}".format(codigo))
        return errores
    numero_planilla = _numero_de_carpeta(fila)
    if not numero_planilla:
        errores.append("carpeta_origen vacio")
    carpeta_codigo = "{:02d}".format(descomponer_codigo(codigo)[0])
    if numero_planilla and numero_planilla != carpeta_codigo:
        errores.append(
            "carpeta_origen {} no coincide con la carpeta {} del codigo {}".format(
                numero_planilla, carpeta_codigo, codigo
            )
        )
    if not str(fila.get("archivo_original", "") or "").strip():
        errores.append("archivo_original vacio")
    if not _entradas(mapa):
        errores.append("mapa de carpetas (hoja Listas) no disponible")
    elif not carpeta_real(mapa, fila.get("carpeta_origen", "")):
        errores.append(
            "la carpeta {} (carpeta_origen) no existe en el mapa de Listas".format(
                numero_planilla
            )
        )
    return errores


def enriquecer_registro(fila, mapa, png=None):
    """Valida un registro y lo enriquece con la trazabilidad completa.

    Devuelve un dict con codigo, carpeta_origen, archivo_original,
    carpeta_real, ruta_cdr (si el mapa lo permite) y las demas columnas de la
    planilla. Lanza ErrorTrazabilidad si el registro es incompleto; en ese
    caso NUNCA se fabrica carpeta_origen ni archivo_original.
    """
    problemas = validar_registro(fila, mapa)
    if problemas:
        raise ErrorTrazabilidad(
            "Registro incompleto (fila {}): {}".format(
                fila.get("_fila", "?"), "; ".join(problemas)
            )
        )
    codigo = str(fila.get("codigo", "") or "").strip()
    carpeta_origen = str(fila.get("carpeta_origen", "") or "").strip()
    archivo_original = str(fila.get("archivo_original", "") or "").strip()
    carpeta = carpeta_real(mapa, carpeta_origen)
    registro = {
        "codigo": codigo,
        "carpeta_origen": carpeta_origen,
        "archivo_original": archivo_original,
        "carpeta_real": carpeta,
    }
    if png:
        registro["png"] = png
    if carpeta:
        registro["ruta_cdr"] = os.path.join(carpeta, archivo_original)
    for columna in COLUMNAS_PLANILLA_DISENO:
        registro[columna] = str(fila.get(columna, "") or "").strip()
    return registro


def filas_por_codigo(filas):
    """Indice codigo -> fila de la planilla y lista de codigos duplicados."""
    indice = {}
    duplicados = []
    for fila in filas:
        codigo = str(fila.get("codigo", "") or "").strip()
        if not codigo:
            continue
        if codigo in indice:
            duplicados.append(codigo)
        indice[codigo] = fila
    return indice, sorted(set(duplicados))


def localizar_cdr(fila, mapa):
    """Reconstruye la ubicacion del CDR original desde un registro completo.

    Devuelve {"codigo", "carpeta_origen", "archivo_original", "carpeta_real",
    "ruta_cdr"}. Lanza ErrorTrazabilidad si falta algo. Esta funcion NO accede
    ni modifica el CDR: solo calcula su ruta a partir del contrato de datos.
    """
    problemas = validar_registro(fila, mapa)
    if problemas:
        raise ErrorTrazabilidad("; ".join(problemas))
    codigo = str(fila.get("codigo", "") or "").strip()
    carpeta_origen = str(fila.get("carpeta_origen", "") or "").strip()
    archivo_original = str(fila.get("archivo_original", "") or "").strip()
    carpeta = carpeta_real(mapa, carpeta_origen)
    if not carpeta:
        raise ErrorTrazabilidad(
            "La carpeta {} (carpeta_origen) no tiene carpeta real en el mapa de Listas.".format(
                _normalizar_numero(carpeta_origen)
            )
        )
    return {
        "codigo": codigo,
        "carpeta_origen": carpeta_origen,
        "archivo_original": archivo_original,
        "carpeta_real": carpeta,
        "ruta_cdr": os.path.join(carpeta, archivo_original),
    }


def localizar_cdr_por_codigo(codigo, filas, mapa):
    """Localiza el CDR de un codigo buscando su fila en la planilla."""
    codigo = str(codigo or "").strip()
    for fila in filas:
        if str(fila.get("codigo", "") or "").strip() == codigo:
            return localizar_cdr(fila, mapa)
    raise ErrorTrazabilidad("No hay fila en la planilla para el codigo {}".format(codigo))